from docx import Document
from docxtpl import DocxTemplate
import openpyxl
from django.http import HttpResponse, JsonResponse, FileResponse
from rest_framework.decorators import api_view
from .models import Student, Payment
import io, os, zipfile, uuid
from django.conf import settings
from datetime import datetime
from django.shortcuts import render, redirect
from django.contrib.auth import authenticate, login, logout
from django.contrib.auth.models import User
from django.contrib import messages
from django.contrib.auth.decorators import login_required
from .serializers import StudentSerializer
from rest_framework import viewsets
from rest_framework.views import APIView
from rest_framework.response import Response
from rest_framework import status
from django.views.decorators.csrf import csrf_exempt
from rest_framework.authtoken.models import Token





class StudentListCreateAPIView(APIView):
    def get(self, request):
        students = Student.objects.all()
        serializer = StudentSerializer(students, many=True)
        return Response(serializer.data)

    def post(self, request):
        serializer = StudentSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)


class StudentRetrieveUpdateDestroyAPIView(APIView):
    def get_object(self, pk):
        try:
            return Student.objects.get(pk=pk)
        except Student.DoesNotExist:
            return None

    def get(self, request, pk):
        student = self.get_object(pk)
        if not student:
            return Response(status=status.HTTP_404_NOT_FOUND)
        serializer = StudentSerializer(student)
        return Response(serializer.data)

    def put(self, request, pk):
        student = self.get_object(pk)
        if not student:
            return Response(status=status.HTTP_404_NOT_FOUND)
        serializer = StudentSerializer(student, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def delete(self, request, pk):
        student = self.get_object(pk)
        if not student:
            return Response(status=status.HTTP_404_NOT_FOUND)
        student.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)

class StudentViewSet(viewsets.ModelViewSet):
    queryset = Student.objects.all()
    serializer_class = StudentSerializer


class StudentListCreateAPIView(APIView):
    def get(self, request):
        students = Student.objects.all()
        serializer = StudentSerializer(students, many=True)
        return Response(serializer.data)

    def post(self, request):
        serializer = StudentSerializer(data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data, status=status.HTTP_201_CREATED)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)



class StudentRetrieveUpdateDestroyAPIView(APIView):
    def get_object(self, pk):
        try:
            return Student.objects.get(pk=pk)
        except Student.DoesNotExist:
            return None

    def get(self, request, pk):
        student = self.get_object(pk)
        if not student:
            return Response(status=status.HTTP_404_NOT_FOUND)
        serializer = StudentSerializer(student)
        return Response(serializer.data)

    def put(self, request, pk):
        student = self.get_object(pk)
        if not student:
            return Response(status=status.HTTP_404_NOT_FOUND)
        serializer = StudentSerializer(student, data=request.data)
        if serializer.is_valid():
            serializer.save()
            return Response(serializer.data)
        return Response(serializer.errors, status=status.HTTP_400_BAD_REQUEST)

    def delete(self, request, pk):
        student = self.get_object(pk)
        if not student:
            return Response(status=status.HTTP_404_NOT_FOUND)
        student.delete()
        return Response(status=status.HTTP_204_NO_CONTENT)



def format_uzbek_date(date_str):
    try:
        date_obj = datetime.strptime(date_str, "%Y-%m-%d")
        oylar = [
            "yanvar", "fevral", "mart", "aprel", "may", "iyun",
            "iyul", "avgust", "sentabr", "oktabr", "noyabr", "dekabr"
        ]
        return f"{date_obj.year}-yil «{date_obj.day}» {oylar[date_obj.month - 1]}"
    except Exception:
        return ""


@login_required(login_url='login')
def upload_excel(request):
    if request.method == 'POST' and request.FILES.get('file'):
        uploaded_file = request.FILES['file']
        filename = uploaded_file.name.lower()
        course = int(request.POST.get("course", 4))
        start_date = request.POST.get("start_date", "")
        end_date = request.POST.get("end_date", "")

        Student.objects.all().delete()

        try:
            if filename.endswith('.xlsx'):
                wb = openpyxl.load_workbook(uploaded_file)
                sheet = wb.active
                for row in sheet.iter_rows(min_row=2, values_only=True):
                    if not row[1]:
                        continue
                    Student.objects.create(
                        full_name=row[1],
                        group=row[2],
                        company=row[3],
                        company_address=row[4],
                        company_director=row[5],
                        company_phone=row[6],
                        practice_supervisor=row[7],
                        faculty=row[8],
                    )

            elif filename.endswith('.docx'):
                doc = Document(uploaded_file)
                table = doc.tables[0]
                for row in table.rows[1:]:
                    cells = row.cells
                    Student.objects.create(
                        full_name=cells[1].text.strip(),
                        group=cells[2].text.strip(),
                        company=cells[3].text.strip(),
                        company_address=cells[4].text.strip(),
                        company_director=cells[5].text.strip(),
                        company_phone=cells[6].text.strip(),
                        practice_supervisor=cells[7].text.strip(),
                        faculty=cells[8].text.strip(),
                    )
            else:
                return render(request, 'app_excel/upload.html', {
                    'error': 'Faqat .xlsx yoki .docx fayl yuklash mumkin.'
                })

            request.session['uploaded'] = True
            request.session['course'] = course
            request.session['start_date'] = start_date
            request.session['end_date'] = end_date
            return redirect('upload_excel')

        except Exception as e:
            return render(request, 'app_excel/upload.html', {'error': str(e)})

    uploaded = request.session.pop('uploaded', False)
    return render(request, 'app_excel/upload.html', {'uploaded': uploaded})


@login_required(login_url='login')
def export_to_word(request):
    students = Student.objects.all()
    if not students.exists():
        return HttpResponse("Hali hech qanday talaba ma'lumotlari mavjud emas.")

    course = int(request.GET.get("course", 4))
    practice_type = "Ishlab chiqarish amaliyoti" if course == 3 else "Bitiruv oldi amaliyoti"
    student_count = students.count()
    first_student = students.first()

    start_date_str = request.session.get("start_date", "2025-02-17")
    formatted_start = format_uzbek_date(start_date_str)

    doc_path = os.path.join(settings.BASE_DIR, 'app_excel', 'templates', 'app_excel', 'shartnoma_template.docx')
    doc = DocxTemplate(doc_path)

    context = {
        'COMPANY': first_student.company,
        'ADDRESS': first_student.company_address,
        'DIRECTOR': first_student.company_director,
        'PHONE': first_student.company_phone,
        'SUPERVISOR': first_student.practice_supervisor,
        'FACULTY': first_student.faculty,
        'COURSE': course,
        'PRACTICE_TYPE': practice_type,
        'STUDENT_COUNT': student_count,
        'START_DATE': formatted_start,
    }

    doc.render(context)
    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    Student.objects.all().delete()

    response = HttpResponse(
        buffer.read(),
        content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document'
    )
    response['Content-Disposition'] = f'attachment; filename="{first_student.company.replace(" ", "_")}_shartnoma.docx"'
    return response




@login_required(login_url='login')
def export_all_documents_zip(request):
    user_profile = request.user.profile
    narx = 30000  # hujjatlarni yuklab olish narxi

    # Mablag' yetarli emas
    if user_profile.balance < narx:
        return HttpResponse("❌ Hisobingizda yetarli mablag‘ mavjud emas.")

    students = Student.objects.all()
    if not students.exists():
        return HttpResponse("❌ Talabalar ma'lumotlari topilmadi.")

    # Foydalanuvchi kiritgan session ma'lumotlari
    course = int(request.session.get("course", 4))
    start_date_str = request.session.get("start_date", "2025-02-17")
    end_date_str = request.session.get("end_date", "2025-04-26")

    # O'zbekcha oylar
    OY_NOMLARI = {
        "January": "yanvar", "February": "fevral", "March": "mart",
        "April": "aprel", "May": "may", "June": "iyun", "July": "iyul",
        "August": "avgust", "September": "sentabr", "October": "oktabr",
        "November": "noyabr", "December": "dekabr"
    }


    start_date = datetime.strptime(start_date_str, "%Y-%m-%d")
    end_date = datetime.strptime(end_date_str, "%Y-%m-%d")
    formatted_start = f"{start_date.year}-yil «{start_date.day:02}» {OY_NOMLARI[start_date.strftime('%B')]}"
    formatted_end = f"{end_date.year}-yil «{end_date.day:02}» {OY_NOMLARI[end_date.strftime('%B')]}"
    student_count = students.count()

    practice_type = "Ishlab chiqarish amaliyoti" if course == 3 else "Bitiruv oldi amaliyoti"


    base_path = os.path.join(settings.BASE_DIR, 'app_excel', 'templates', 'app_excel')
    shartnoma_tpl = os.path.join(base_path, 'shartnoma_template.docx')
    kundalik_tpl = os.path.join(base_path, 'kundalik_template.docx')
    yollanma_tpl = os.path.join(base_path, 'yollanma_template.docx')


    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zip_file:
        for index, student in enumerate(students, start=1):
            folder = f"Hujjatlar/{index}"
            context = {
                'FULL_NAME': student.full_name,
                'GROUP': student.group,
                'COMPANY': student.company,
                'ADDRESS': student.company_address,
                'DIRECTOR': student.company_director,
                'PHONE': student.company_phone,
                'SUPERVISOR': student.practice_supervisor,
                'FACULTY': student.faculty,
                'COURSE': course,
                'PRACTICE_TYPE': practice_type,
                'STUDENT_COUNT': student_count,
                'START_DATE': formatted_start,
                'END_DATE': formatted_end,
            }


            doc1 = DocxTemplate(shartnoma_tpl)
            doc1.render(context)
            io1 = io.BytesIO()
            doc1.save(io1)
            io1.seek(0)
            zip_file.writestr(f"{folder}/shartnoma.docx", io1.read())


            doc2 = DocxTemplate(kundalik_tpl)
            doc2.render(context)
            io2 = io.BytesIO()
            doc2.save(io2)
            io2.seek(0)
            zip_file.writestr(f"{folder}/kundalik.docx", io2.read())


            doc3 = DocxTemplate(yollanma_tpl)
            doc3.render(context)
            io3 = io.BytesIO()
            doc3.save(io3)
            io3.seek(0)
            zip_file.writestr(f"{folder}/yollanma.docx", io3.read())

    user_profile.balance -= narx
    user_profile.save()

    Student.objects.all().delete()

    zip_buffer.seek(0)
    response = HttpResponse(zip_buffer, content_type='application/zip')
    response['Content-Disposition'] = 'attachment; filename=Hujjatlar.zip'
    return response



@login_required(login_url='login')
def generate_contract_for_company(request, company_name):
    students = Student.objects.filter(company=company_name)
    if not students.exists():
        return HttpResponse(" Bu korxona bo‘yicha talabalar topilmadi.")

    course = int(request.GET.get("course", 4))
    practice_type = "Ishlab chiqarish amaliyoti" if course == 3 else "Bitiruv oldi amaliyoti"

    first_student = students.first()
    student_count = students.count()

    doc = DocxTemplate(os.path.join(
        settings.BASE_DIR, 'app_excel', 'templates', 'app_excel', 'shartnoma_template.docx'))

    context = {
        'COMPANY': first_student.company,
        'ADDRESS': first_student.company_address,
        'DIRECTOR': first_student.company_director,
        'PHONE': first_student.company_phone,
        'SUPERVISOR': first_student.practice_supervisor,
        'FACULTY': first_student.faculty,
        'COURSE': course,
        'PRACTICE_TYPE': practice_type,
        'STUDENT_COUNT': student_count,
    }

    doc.render(context)

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)

    response = HttpResponse(
        buffer.read(), content_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document')
    response['Content-Disposition'] = f'attachment; filename="{company_name.replace(" ", "_")}_shartnoma.docx"'
    return response



def login_view(request):
    if request.method == "POST":
        username = request.POST.get("username")
        password = request.POST.get("password")
        remember = request.POST.get("remember")

        user = authenticate(request, username=username, password=password)
        if user:
            login(request, user)

            if remember:
                request.session.set_expiry(31536000)
            else:
                request.session.set_expiry(0)

            return redirect("upload_excel")
        else:
            messages.error(request, " Login yoki parol noto‘g‘ri.")

    return render(request, "app_excel/login.html")


@csrf_exempt
def register_view(request):
    if request.method == "POST":
        username = request.POST.get("username")
        password1 = request.POST.get("password1")
        password2 = request.POST.get("password2")

        if password1 != password2:
            messages.error(request, "Parollar mos emas!")
        elif User.objects.filter(username=username).exists():
            messages.error(request, "Bu login allaqachon mavjud!")
        else:
            user = User.objects.create_user(username=username, password=password1)

            messages.success(request, "Muvaffaqiyatli ro‘yxatdan o‘tildi.")
            return redirect("login")

    return render(request, "app_excel/register.html")


def logout_view(request):
    logout(request)
    return redirect("login")


@login_required(login_url='login')
def pay_view(request):
    user = request.user
    amount = 30000
    merchant_id = '123456'
    service_id = '654321'
    callback_url = 'http://127.0.0.1:8000/payment/callback/'

    merchant_trans_id = str(uuid.uuid4())

    payment = Payment.objects.create(
        user=user,
        amount=amount,
        merchant_trans_id=merchant_trans_id,
        status='pending'  # optional
    )

    context = {
        'payment': payment,
        'merchant_id': merchant_id,
        'service_id': service_id,
        'amount': amount,
        'merchant_trans_id': merchant_trans_id,
        'callback_url': callback_url,
    }

    return render(request, 'app_excel/pay.html', {
        'payment': payment,
        'merchant_id': 'YOUR_MERCHANT_ID',
        'service_id': 'YOUR_SERVICE_ID',
        'amount': amount,
    })


@csrf_exempt
def click_prepare(request):
    return JsonResponse({'error': 0, 'error_note': 'Success'})


@csrf_exempt
def click_result(request):
    merchant_trans_id = request.POST.get("merchant_trans_id")
    click_trans_id = request.POST.get("click_trans_id")

    try:
        payment = Payment.objects.get(merchant_trans_id=merchant_trans_id)

        if not payment.paid:
            payment.paid = True
            payment.click_trans_id = click_trans_id
            payment.save()

            zip_buffer = generate_documents_zip()

            response = HttpResponse(zip_buffer, content_type='application/zip')
            response['Content-Disposition'] = f'attachment; filename=hujjatlar.zip'
            return response

        return JsonResponse({'error': 0, 'error_note': 'To‘lov avval amalga oshirilgan'})
    except Payment.DoesNotExist:
        return JsonResponse({'error': -5, 'error_note': 'To‘lov topilmadi'})



def generate_documents_zip():
    students = Student.objects.all()
    course = 4
    practice_type = "Bitiruv oldi amaliyoti"

    start_date = datetime(2025, 2, 17)
    end_date = datetime(2025, 4, 26)

    formatted_start = f"{start_date.year}-yil «{start_date.day}» fevral"
    formatted_end = f"{end_date.year}-yil «{end_date.day}» aprel"

    shartnoma_path = os.path.join(settings.BASE_DIR, 'app_excel/templates/app_excel/shartnoma_template.docx')
    kundalik_path = os.path.join(settings.BASE_DIR, 'app_excel/templates/app_excel/kundalik_template.docx')
    yollanma_path = os.path.join(settings.BASE_DIR, 'app_excel/templates/app_excel/yollanma_template.docx')

    zip_buffer = io.BytesIO()
    with zipfile.ZipFile(zip_buffer, "w") as zip_file:
        for index, student in enumerate(students, start=1):
            context = {
                'FULL_NAME': student.full_name,
                'GROUP': student.group,
                'COMPANY': student.company,
                'ADDRESS': student.company_address,
                'DIRECTOR': student.company_director,
                'PHONE': student.company_phone,
                'SUPERVISOR': student.practice_supervisor,
                'FACULTY': student.faculty,
                'COURSE': course,
                'PRACTICE_TYPE': practice_type,
                'START_DATE': formatted_start,
                'END_DATE': formatted_end,
            }

            folder = f"Hujjatlar/{index}"

            for tpl_path, name in [(shartnoma_path, "shartnoma"), (kundalik_path, "kundalik"), (yollanma_path, "yollanma")]:
                tpl = DocxTemplate(tpl_path)
                tpl.render(context)
                io_doc = io.BytesIO()
                tpl.save(io_doc)
                io_doc.seek(0)
                zip_file.writestr(f"{folder}/{name}.docx", io_doc.read())

    zip_buffer.seek(0)
    return zip_buffer


@csrf_exempt
def payment_callback(request):
    if request.method == "POST" or request.method == "GET":

        request.session['paid'] = True

        return redirect('export_all_documents_zip')

    return HttpResponse("To‘lov bekor qilindi yoki noto‘g‘ri so‘rov!", status=400)



@login_required
def top_up_balance(request):
    if request.method == "POST":
        amount = int(request.POST.get("amount", 0))
        profile = request.user.profile
        profile.balance += amount
        profile.save()
        return redirect('top_up_balance')

    return render(request, 'balance.html')


@login_required
def balance_view(request):
    return render(request, 'balance.html')



@login_required(login_url='login')
def profile_view(request):
    return render(request, 'app_excel/profile.html', {
        'user': request.user,
        'profile': request.user.profile
    })


@login_required(login_url='login')
def account_settings_view(request):
    return render(request, 'app_excel/account_settings.html', {
        'user': request.user
    })



@api_view(['POST'])
def custom_login_api(request):
    username = request.data.get("username")
    password = request.data.get("password")

    user = authenticate(username=username, password=password)
    if user:
        token, created = Token.objects.get_or_create(user=user)
        return Response({"token": token.key, "username": user.username})
    return Response({"error": "Login yoki parol noto‘g‘ri"}, status=400)



def custom_401_view(request, exception=None):
    return render(request, 'errors/401.html', status=401)

def custom_404_view(request, exception=None):
    return render(request, 'errors/404.html', status=404)

def custom_500_view(request):
    return render(request, 'errors/500.html', status=500)

def csrf_failure(request, reason=""):
    return render(request, 'errors/403_csrf.html', status=403)







